from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path("/Users/utente/Documents/semantic-search - POC/open-reply-poc/deliverables/POC_Semantic_Search_Guida_Tecnica_QA.docx")

GREEN = "28B779"
DARK = "111111"
MID = "5F6368"
LIGHT = "E7ECE9"
PALE = "F4F8F6"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_row(row, allow_split=False):
    tr_pr = row._tr.get_or_add_trPr()
    if not allow_split:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_bottom_border(paragraph, color=GREEN, size=10):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.7)
section.footer_distance = Cm(0.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for style_name, size, before, after in (
    ("Title", 29, 0, 14),
    ("Heading 1", 21, 18, 8),
    ("Heading 2", 15, 14, 6),
    ("Heading 3", 11.5, 9, 3),
):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Heading 3" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(DARK)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

if "Question" not in styles:
    q_style = styles.add_style("Question", WD_STYLE_TYPE.PARAGRAPH)
else:
    q_style = styles["Question"]
q_style.font.name = "Aptos"
q_style.font.size = Pt(11)
q_style.font.bold = True
q_style.font.color.rgb = RGBColor.from_string(DARK)
q_style.paragraph_format.space_before = Pt(9)
q_style.paragraph_format.space_after = Pt(2)
q_style.paragraph_format.keep_with_next = True

if "Answer" not in styles:
    a_style = styles.add_style("Answer", WD_STYLE_TYPE.PARAGRAPH)
else:
    a_style = styles["Answer"]
a_style.font.name = "Aptos"
a_style.font.size = Pt(10.2)
a_style.font.color.rgb = RGBColor.from_string(DARK)
a_style.paragraph_format.space_after = Pt(7)
a_style.paragraph_format.line_spacing = 1.1

header = section.header
hp = header.paragraphs[0]
hp.text = "ON DEVICE SEMANTIC SEARCH  |  GUIDA TECNICA"
hp.style = styles["Caption"]
hp.runs[0].font.name = "Aptos"
hp.runs[0].font.size = Pt(8)
hp.runs[0].font.bold = True
hp.runs[0].font.color.rgb = RGBColor.from_string(MID)

footer = section.footer
fp = footer.paragraphs[0]
fp.add_run("Open Reply  |  Documento di preparazione al talk     ")
fp.runs[0].font.name = "Aptos"
fp.runs[0].font.size = Pt(8)
fp.runs[0].font.color.rgb = RGBColor.from_string(MID)
add_page_number(fp)


title = doc.add_paragraph(style="Title")
title.add_run("POC On Device\nSemantic Search")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_bottom_border(title)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_before = Pt(4)
subtitle.paragraph_format.space_after = Pt(18)
r = subtitle.add_run("Guida tecnica e domande per il talk")
r.font.name = "Aptos Display"
r.font.size = Pt(16)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(GREEN)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run(
    "Questo documento riassume il funzionamento della Proof of Concept e raccoglie le principali "
    "domande tecniche che potrebbero emergere durante una presentazione. Le risposte distinguono "
    "le funzionalità già realizzate dalle possibili evoluzioni verso una soluzione industriale."
)
r.font.size = Pt(11)

meta = doc.add_table(rows=3, cols=2)
meta.autofit = False
meta.columns[0].width = Cm(4.2)
meta.columns[1].width = Cm(11.7)
meta.style = "Table Grid"
meta_data = [
    ("Knowledge base finale", "1.709 prodotti: 709 iniziali e 1.000 aggiuntivi"),
    ("Modello selezionato", "Xenova/multilingual-e5-small, quantizzazione q4, 384 dimensioni"),
    ("Esecuzione", "Generazione degli embedding, ricerca e ranking nel browser"),
]
for idx, (label, value) in enumerate(meta_data):
    row = meta.rows[idx]
    set_repeat_table_row(row)
    for cell in row.cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(row.cells[0], PALE)
    pr = row.cells[0].paragraphs[0]
    pr.paragraph_format.space_after = Pt(0)
    rr = pr.add_run(label)
    rr.bold = True
    rr.font.color.rgb = RGBColor.from_string(DARK)
    pr = row.cells[1].paragraphs[0]
    pr.paragraph_format.space_after = Pt(0)
    pr.add_run(value)

doc.add_paragraph()
h = doc.add_paragraph("Contenuti", style="Heading 2")
for item in (
    "1. Flusso tecnico della POC",
    "2. Architettura",
    "3. Modello ed embedding",
    "4. Ricerca e ranking",
    "5. Filtri automatici",
    "6. Prestazioni",
    "7. Dati, qualità e limiti",
    "8. Domande critiche",
):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(item)

doc.add_page_break()

doc.add_heading("1. Flusso tecnico della POC", level=1)

doc.add_paragraph(
    "La POC dimostra che una ricerca semantica può essere eseguita direttamente nel frontend. "
    "Il browser carica un modello di embedding tramite Transformers.js, trasforma query e prodotti "
    "in vettori numerici confrontabili e calcola il ranking senza inviare ogni ricerca a un servizio AI esterno."
)

doc.add_heading("Preparazione della knowledge base", level=2)
doc.add_paragraph(
    "I prodotti arrivano come dati strutturati. Il codice applicativo seleziona e normalizza gli attributi "
    "rilevanti, quindi costruisce per ogni prodotto una descrizione testuale uniforme. Questa descrizione "
    "non viene scritta dal modello: è prodotta in modo deterministico dal nostro codice, così prodotti diversi "
    "vengono rappresentati secondo lo stesso schema."
)

doc.add_heading("Generazione e riuso degli embedding", level=2)
doc.add_paragraph(
    "Al primo caricamento utile, il modello Xenova/multilingual-e5-small viene inizializzato nel browser. "
    "Le descrizioni dei prodotti, precedute dal prefisso “passage:”, vengono tokenizzate e trasformate in "
    "embedding normalizzati di 384 valori. Gli embedding sono associati agli identificativi dei prodotti e "
    "persistiti con le informazioni di versione del modello, in modo da poterli recuperare negli accessi successivi."
)
doc.add_paragraph(
    "Quando l’applicazione viene riaperta, la copia in memoria viene ricostruita recuperando gli embedding già "
    "disponibili. Se il modello o la versione della rappresentazione testuale cambiano, gli elementi non più "
    "compatibili vengono rigenerati. I file del modello sono inoltre memorizzati nella cache del browser."
)

doc.add_heading("Esecuzione della ricerca", level=2)
doc.add_paragraph(
    "La query dell’utente è già testo in linguaggio naturale e non richiede la costruzione della descrizione "
    "usata per i prodotti. Viene preceduta dal prefisso “query:” e trasformata dallo stesso modello in un embedding. "
    "Il sistema confronta quindi il vettore della query con tutti i vettori dei prodotti attraverso la similarità "
    "coseno. Al punteggio semantico possono essere sommati boost derivati da regole di business configurabili; "
    "il risultato finale determina l’ordinamento e i Top Match."
)

doc.add_heading("Esperienza AI-driven", level=2)
doc.add_paragraph(
    "L’interfaccia rende immediatamente visibili i tre risultati migliori, il relativo score e i criteri che hanno "
    "contribuito al ranking. In parallelo, alcune informazioni riconoscibili nella query possono precompilare i "
    "filtri della ricerca avanzata. Nella versione attuale questo secondo comportamento è deterministico e basato "
    "su regole; l’adozione di un ulteriore modello locale per il query understanding è una possibile evoluzione."
)

flow = doc.add_table(rows=1, cols=7)
flow.autofit = False
widths = [2.4, 2.1, 2.4, 2.0, 2.4, 2.0, 2.4]
labels = [
    "Prodotti\nnormalizzati",
    "Embedding\nprodotti",
    "Indice\nin memoria",
    "Query →\nembedding",
    "Similarità\ncoseno",
    "Ranking\nfinale",
    "Top Match\ne filtri",
]
for i, cell in enumerate(flow.rows[0].cells):
    cell.width = Cm(widths[i])
    set_cell_margins(cell, top=90, bottom=90, start=55, end=55)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, PALE if i % 2 == 0 else "FFFFFF")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(labels[i])
    run.bold = True
    run.font.size = Pt(8.2)
set_repeat_table_row(flow.rows[0])

doc.add_page_break()
doc.add_heading("2. Domande e risposte", level=1)
doc.add_paragraph(
    "Le risposte seguenti sono formulate per essere utilizzate anche oralmente. Dove opportuno, viene indicato "
    "il confine tra la POC attuale e una possibile industrializzazione."
)


qa_sections = [
    ("Architettura", [
        ("1. Cosa significa realmente “on-device”?",
         "Significa che l’inferenza necessaria a trasformare la query in embedding e il confronto con gli embedding dei prodotti avvengono nel browser dell’utente. Il backend continua a distribuire il dataset e a supportare la persistenza. Quindi la POC è on-device per la parte AI e di ricerca, ma non è un’applicazione completamente offline."),
        ("2. Qual è il ruolo di Transformers.js?",
         "Transformers.js consente di scaricare, inizializzare ed eseguire nel browser un modello compatibile con il formato ONNX. Nel nostro caso espone una pipeline di feature extraction: riceve il testo tokenizzato, esegue il modello e restituisce un vettore normalizzato utilizzabile per la ricerca semantica."),
        ("3. Dove vengono memorizzati il modello e gli embedding?",
         "I file del modello vengono memorizzati dalla cache del browser. Gli embedding dei prodotti vengono persistiti lato backend insieme all’identificativo del prodotto e alla versione del modello; quando la pagina è attiva, vengono inoltre caricati in un indice in memoria, cioè una struttura JavaScript disponibile nella RAM del browser."),
        ("4. Gli embedding vengono rigenerati a ogni accesso?",
         "No. All’avvio il sistema verifica se esiste un embedding compatibile per ciascun prodotto. Se modello, versione e rappresentazione testuale coincidono, l’embedding viene riutilizzato. Vengono rigenerati soltanto gli elementi mancanti o non più compatibili."),
        ("5. Perché è stato usato SQLite e non un database vettoriale?",
         "Per una POC con 1.709 prodotti, SQLite è sufficiente per persistere vettori e metadati senza introdurre ulteriore infrastruttura. La ricerca non viene eseguita da SQLite: i vettori vengono caricati in memoria e confrontati nel browser. Un database vettoriale o un indice ANN diventerebbero opportuni con cataloghi molto più grandi o requisiti di ricerca server-side."),
    ]),
    ("Modello ed embedding", [
        ("6. Come viene trasformato un prodotto in un embedding?",
         "Il codice estrae i campi rilevanti, normalizza valori e categorie e costruisce una descrizione testuale uniforme. Il tokenizer suddivide il testo in token e li converte in identificativi numerici. Il modello elabora questi token nel loro contesto e produce rappresentazioni interne; con mean pooling e normalizzazione si ottiene un singolo vettore di 384 dimensioni."),
        ("7. È il modello a costruire la descrizione del prodotto?",
         "No. La descrizione viene costruita dal codice applicativo con un template deterministico. Il modello interviene dopo e trasforma quella descrizione in un embedding. Questa separazione rende la rappresentazione ripetibile e controllabile."),
        ("8. Come riconosce sinonimi e contesto?",
         "La conoscenza linguistica deriva dall’addestramento del modello su grandi quantità di coppie e frasi semanticamente correlate. Durante l’inferenza i token non vengono interpretati isolatamente: i meccanismi di attenzione producono rappresentazioni contestuali, per cui termini diversi possono risultare vicini quando esprimono concetti simili."),
        ("9. Perché un modello di embedding invece di un LLM?",
         "Il problema principale è il retrieval, cioè confrontare rapidamente una query con molti elementi. Un modello di embedding è più leggero, produce vettori riutilizzabili, ha latenza e requisiti inferiori ed è adatto al calcolo massivo della similarità. Un LLM è più indicato per comprendere o riscrivere richieste complesse e generare spiegazioni; può essere aggiunto in una fase separata, senza sostituire il motore di retrieval."),
        ("10. Perché vengono usati i prefissi “query:” e “passage:”?",
         "Il modello E5 è stato addestrato distinguendo testi di ricerca e documenti da recuperare. I prefissi indicano il ruolo del testo e permettono di utilizzare il modello come previsto dal suo addestramento: “query:” per la richiesta e “passage:” per la descrizione del prodotto."),
        ("11. Perché è stato scelto multilingual-e5-small?",
         "Nel confronto svolto nella POC ha offerto il miglior equilibrio tra qualità, dimensione e velocità. Supporta più lingue, produce vettori compatti da 384 dimensioni ed è abbastanza leggero per l’esecuzione nel browser. Sul benchmark manuale di nove query ha ottenuto otto risultati giudicati buoni e uno parziale, senza casi deboli."),
    ]),
    ("Ricerca e ranking", [
        ("12. Come viene calcolata la similarità?",
         "Per ogni prodotto viene calcolata la similarità coseno tra il suo embedding e quello della query. La misura confronta la direzione dei due vettori: più sono allineati, maggiore è la vicinanza semantica. Poiché gli embedding sono normalizzati, il calcolo è efficiente e stabile."),
        ("13. Perché è stata scelta la similarità coseno?",
         "È la metrica normalmente associata ai modelli di sentence embedding e misura l’orientamento semantico senza essere influenzata dalla lunghezza del vettore. Per il modello adottato è una scelta coerente con il modo in cui gli embedding vengono addestrati e normalizzati."),
        ("14. Lo score mostrato dipende soltanto dal modello?",
         "No. Il sistema calcola uno score semantico e può aggiungere un business boost derivato da regole esplicite. Lo score finale è quindi la somma dei due contributi. Nei log vengono mantenuti separati semanticScore, businessBoost, finalScore e matchedRules, così il ranking resta ispezionabile."),
        ("15. Le regole di business non rischiano di alterare troppo i risultati?",
         "Sì, se configurate con pesi eccessivi. Per questo devono essere trasparenti, misurabili e validate su query rappresentative. Il vantaggio dell’approccio ibrido è poter riflettere vincoli applicativi; il rischio è mascherare la qualità semantica. I due contributi vanno quindi monitorati separatamente."),
        ("16. Perché vengono evidenziati proprio tre Top Match?",
         "Tre risultati consentono di dare una risposta immediata senza nascondere il resto del catalogo. È una scelta di interfaccia, non un limite del motore: la lista completa resta disponibile e il numero di elementi evidenziati può essere configurato."),
        ("17. La ricerca semantica sostituisce la ricerca esatta?",
         "No. Identificativi, codici e nomi precisi devono beneficiare di corrispondenze esatte o di un ranking ibrido. La ricerca semantica è particolarmente utile per bisogni descrittivi e concetti approssimativi; una soluzione completa combina ricerca esatta, filtri e similarità semantica."),
    ]),
    ("Filtri automatici", [
        ("18. È il modello E5 a popolare automaticamente i filtri?",
         "No. Nella POC attuale il popolamento dei filtri è deterministico. Il testo viene normalizzato e confrontato con espressioni regolari, parole chiave e gruppi di sinonimi configurati nel codice. Sono gestiti valori espliciti, categorie, valuta, attributi booleani e alcune forme di negazione."),
        ("19. Possiamo comunque definire l’interfaccia “AI-driven”?",
         "Sì, perché l’esperienza complessiva è guidata dal ranking semantico e presenta Top Match, score e criteri riconosciuti. È però corretto precisare che l’autocompilazione dei filtri è oggi rule-based. Un secondo modello locale per il query understanding è una possibile evoluzione, non una funzionalità già presente."),
        ("20. Cosa succede con un sinonimo non previsto dalle regole dei filtri?",
         "Il retrieval semantico può comunque recuperare prodotti coerenti, perché il modello riconosce relazioni linguistiche apprese. Il filtro, invece, potrebbe non essere precompilato. È il limite principale dell’approccio deterministico e il motivo per cui è stata valutata l’aggiunta futura di un modello dedicato alla classificazione degli intenti."),
    ]),
    ("Prestazioni", [
        ("21. Come sono stati misurati i tempi della query?",
         "Il codice usa performance.now() prima e dopo l’intera funzione di ricerca. La misura comprende la generazione dell’embedding della query, il confronto coseno con tutti i prodotti, l’applicazione dei boost e l’ordinamento. Non comprende il rendering dell’interfaccia, il download iniziale del modello o la generazione iniziale dell’indice."),
        ("22. Cosa significa “query warm”?",
         "È una ricerca eseguita quando il modello è già caricato, gli embedding dei prodotti sono già disponibili in memoria e il motore ha già completato l’inizializzazione. Misura quindi il comportamento normale dopo il primo avvio, senza includere i costi di preparazione."),
        ("23. I 21 millisecondi sono un valore garantito?",
         "No. È la media osservata nella macchina usata per la POC sulle nove query di benchmark con il modello E5 e il catalogo disponibile. Il valore dipende da hardware, browser, numero di prodotti e carico del dispositivo; va presentato come misura sperimentale, non come SLA."),
        ("24. Perché il primo caricamento è più lento?",
         "La prima esecuzione può richiedere il download e l’inizializzazione del modello, il controllo della cache e la generazione degli embedding mancanti. Nel benchmark su 1.509 prodotti, la generazione completa con E5 ha richiesto circa 6 minuti e 16 secondi. Il catalogo finale è poi cresciuto a 1.709 elementi, quindi i due numeri non descrivono lo stesso dataset."),
        ("25. Come viene evitato il blocco dell’interfaccia?",
         "L’indicizzazione viene suddivisa in batch e il controllo viene periodicamente restituito al browser, così la pagina può continuare a reagire. Per carichi più elevati o per un secondo modello locale, l’evoluzione naturale è spostare l’elaborazione in un Web Worker, separandola dal thread che gestisce l’interfaccia."),
    ]),
    ("Dati, qualità e limiti", [
        ("26. Il modello conosce già i prodotti del catalogo?",
         "No. Il modello conosce strutture linguistiche generali, non il catalogo specifico. Sono la knowledge base e le descrizioni costruite dal codice a fornire le informazioni sui prodotti. Il modello le traduce in uno spazio vettoriale confrontabile."),
        ("27. Cosa accade se i dati dei prodotti sono incompleti?",
         "La qualità del retrieval diminuisce: un embedding non può rappresentare caratteristiche che non sono presenti nella descrizione. La completezza, la coerenza e la normalizzazione della knowledge base sono quindi parte integrante della qualità della soluzione."),
        ("28. Come è stata valutata la qualità dei modelli?",
         "Sono state usate nove query rappresentative. Per ciascuna query sono stati analizzati manualmente i primi tre risultati e classificati come buoni, parziali o deboli. È un benchmark utile per la POC, ma non sostituisce una valutazione più ampia con giudizi di dominio, metriche come Precision@K o nDCG e un insieme di query versionato."),
        ("29. Perché compaiono sia 1.509 sia 1.709 prodotti?",
         "Il benchmark di generazione e confronto tra modelli è stato eseguito quando il dataset conteneva 1.509 prodotti. La knowledge base finale è poi arrivata a 1.709 elementi. Nel talk bisogna associare sempre ogni misura alla versione del dataset utilizzata."),
        ("30. La soluzione è già pronta per la produzione?",
         "No. La POC dimostra fattibilità tecnica e valore dell’esperienza utente. Per un’industrializzazione servono benchmark più ampi, gestione robusta degli errori, osservabilità, accessibilità, policy di aggiornamento del modello, sicurezza, test su più dispositivi e una strategia di scalabilità per cataloghi più grandi."),
    ]),
    ("Domande critiche", [
        ("31. Perché eseguire questa logica nel frontend invece che su un server?",
         "Il frontend riduce la latenza dopo l’inizializzazione, evita una chiamata AI per ogni query, mantiene la richiesta sul dispositivo e dimostra che il retrieval può essere integrato in applicazioni esistenti con infrastruttura ridotta. Il compromesso è trasferire download, memoria e calcolo al dispositivo dell’utente."),
        ("32. La query dell’utente viene inviata a terze parti?",
         "Durante la ricerca no: la query viene elaborata localmente dal modello già caricato. Al primo utilizzo, i file del modello sono scaricati dall’hosting configurato, attualmente Hugging Face, a meno di scegliere il self-hosting. Questa distinzione va esplicitata quando si parla di privacy e on-device."),
        ("33. Cosa succede se il modello non si carica?",
         "La POC prevede un fallback mock utile allo sviluppo. In una versione industriale sarebbe preferibile mostrare chiaramente l’indisponibilità della ricerca semantica e mantenere funzionanti ricerca tradizionale e filtri, evitando di presentare risultati simulati come risultati reali."),
        ("34. Qual è oggi il limite tecnico principale?",
         "L’indice in memoria è una semplice collezione JavaScript e ogni query confronta linearmente tutti i vettori. Con 1.709 prodotti è una soluzione adeguata; con cataloghi molto più grandi sarebbe necessario valutare un indice ANN, un motore vettoriale o una ricerca server-side. Anche il query understanding dei filtri è ancora basato su regole."),
    ]),
]

for section_title, questions in qa_sections:
    doc.add_heading(section_title, level=2)
    for question, answer in questions:
        qp = doc.add_paragraph(question, style="Question")
        ap = doc.add_paragraph(answer, style="Answer")

doc.add_page_break()
doc.add_heading("Risposta prudente quando manca una misura", level=1)
quote = doc.add_paragraph()
quote.paragraph_format.left_indent = Cm(0.8)
quote.paragraph_format.right_indent = Cm(0.8)
quote.paragraph_format.space_before = Pt(10)
quote.paragraph_format.space_after = Pt(10)
r = quote.add_run(
    "“Non abbiamo ancora un dato sufficientemente validato per dare un numero definitivo. "
    "Nella POC abbiamo verificato la fattibilità; quel punto rientra tra le misurazioni necessarie "
    "prima di un’eventuale industrializzazione.”"
)
r.italic = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor.from_string(DARK)

doc.add_paragraph(
    "Questa formulazione evita stime non dimostrate e mantiene il focus sul risultato effettivamente raggiunto: "
    "la fattibilità della ricerca semantica nel browser e il valore dell’interfaccia AI-driven."
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
